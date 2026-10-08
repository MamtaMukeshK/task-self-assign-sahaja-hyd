"""Builds docs/followup/sample-followup-sheet.xlsx: the Follow-up Program sheet filled with MADE-UP demo data.

The tabs are produced by the real followup/Code.gs (run on the simulated sheet in test/harness.js), so column titles,
generated dates and the Volunteers cell format are exactly what the live sheet will have. Same data as the demo videos.
Run from the repo root:  python3 tools/make_followup_sample.py   (needs node and openpyxl)
"""
import datetime, json, pathlib, subprocess
from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = pathlib.Path(__file__).resolve().parent.parent
JS = r"""
const { load, makeSheet } = require('./test/harness');
const gs0 = load([], new Date(), 'followup/Code.gs');
const PH = gs0.PLAN_COLS.map(c => c[1]), SH = gs0.SLOT_COLS.map(c => c[1]);
const line = o => PH.map(h => o[h] == null ? '' : o[h]);
const map = q => 'https://maps.google.com/?q=' + encodeURIComponent(q + ', Hyderabad');
// One line per frequency, so every kind of line has an example. Places, people and phone numbers are made up.
const plan = makeSheet('Program plan', 1, [PH,
  line({ Day: 'Saturday', Start: '6:30 PM', End: '7:30 PM', Frequency: 'Weekly', Centre: 'Ameerpet Centre', Address: 'Road 3, Ameerpet',
    'Google map': map('Ameerpet'), Places: '4', From: '2026-10-01', Until: '2026-12-31', Contact: 'Lakshmi 9000000009' }),
  line({ Day: 'Every day', Start: '7:00 AM', End: '7:45 AM', Frequency: 'Daily', Centre: 'Kukatpally Centre', Address: 'KPHB Colony',
    'Google map': map('KPHB Colony'), Places: '2', From: '2026-10-12', Until: '2026-10-25', Contact: 'Lakshmi 9000000009' }),
  line({ Start: '6:30 PM', End: '7:30 PM', Frequency: 'Daily, weekdays only (Mon-Fri)', Centre: 'Dilsukhnagar Centre', Address: 'Near the bus stand',
    'Google map': map('Dilsukhnagar'), Places: '2', From: '2026-10-12', Until: '2026-11-30', Contact: 'Suresh 9000000008' }),
  line({ Day: 'Sunday', Start: '10:00 AM', End: '11:00 AM', Frequency: 'Every 2 weeks', Centre: 'Secunderabad Centre', Address: 'SP Road',
    'Google map': map('Secunderabad'), Places: '3', From: '2026-10-04', Contact: 'Suresh 9000000008', Notes: 'No Until date: keeps 12 weeks ahead' }),
  line({ Day: 'Saturday', Start: '4:00 PM', End: '5:30 PM', Frequency: 'Monthly, same weekday', 'Week of month': '2nd', Centre: 'Gachibowli Centre',
    'Google map': map('Gachibowli'), Places: '5', From: '2026-10-01', Until: '2027-03-31', Contact: 'Lakshmi 9000000009' }),
  line({ Start: '5:00 PM', End: '6:00 PM', Frequency: 'Monthly, same date', 'Day of month': '15', Centre: 'Madhapur Centre',
    'Google map': map('Madhapur'), Places: '2', From: '2026-10-01', Until: '2027-03-31', Contact: 'Suresh 9000000008' }),
  line({ Start: '11:00 AM', End: '12:30 PM', Frequency: 'One-off', Centre: 'Begumpet Centre', 'Google map': map('Begumpet'), Places: '6',
    From: '2026-11-14', Contact: 'Lakshmi 9000000009', Notes: "Children's Day programme" })]);
const slots = makeSheet('Slots', 2, [SH]);
const speakers = makeSheet('Speakers', 3, [['Sr. No.', 'Speaker', 'Mobile'], ['1', 'Asha Rao', '9000000001'], ['2', 'Ravi Kumar', '9000000002'],
  ['3', 'Meena Iyer', ''], ['4', 'Gita Sharma', '9000000004'], ['5', 'Hari Prasad', '9000000005']]);
const gs = load([plan, slots, speakers], new Date(Date.UTC(2026, 9, 8, 5, 0)), 'followup/Code.gs');
gs.generateSlots();
gs.registerSlots(['P1-20261010'], 4, 'Asha Rao', '9000000001', 'Ravi Kumar 9000000002');
gs.registerSlots(['P3-20261012'], 1, 'Gita Sharma', '9000000004');
gs.registerSlots(['P1-20261017'], 1, 'Hari Prasad', '9000000005');
slots.grid.find(r => r[0] === 'P1-20261024')[SH.indexOf('Status')] = 'Cancelled';
gs.generateSlots(); // refreshes the Update report
const report = gs.SpreadsheetApp.getActiveSpreadsheet().getSheetByName('Update report');
console.log(JSON.stringify({ plan: plan.grid, slots: slots.grid, speakers: speakers.grid, report: report.grid,
  planNotes: gs0.PLAN_COLS.map(c => c[2]), slotNotes: gs0.SLOT_COLS.map(c => c[2]), freqs: gs0.FREQUENCIES, days: gs0.DAY_NAMES }));
"""
data = json.loads(subprocess.run(['node', '-e', JS], cwd=ROOT, capture_output=True, text=True, check=True).stdout)

BOLD, BAND = Font(bold=True), PatternFill('solid', fgColor='E1F0FB')
def cell_value(title, v):
    """Dates and times as real spreadsheet values (as the live sheet holds them); everything else as text or numbers."""
    s = str(v)
    if title in ('Date', 'From', 'Until') and len(s) == 10 and s[4] == '-':
        return datetime.date.fromisoformat(s), 'ddd d mmm yyyy' if title == 'Date' else 'd mmm yyyy'
    if title in ('Start', 'End') and s:
        if ':' in s and len(s) == 5:
            return datetime.time(int(s[:2]), int(s[3:])), 'h:mm AM/PM'
        t = datetime.datetime.strptime(s.replace(' ', ''), '%I:%M%p').time()
        return t, 'h:mm AM/PM'
    if title in ('Places', 'Still needed', 'Sr. No.') and s.isdigit():
        return int(s), None
    return v, None

def add_tab(wb, name, grid, notes=None, widths=None):
    ws = wb.create_sheet(name)
    heads = grid[0]
    for r, row in enumerate(grid, 1):
        for c, v in enumerate(row, 1):
            if r == 1:
                cell = ws.cell(r, c, v); cell.font = BOLD; cell.fill = BAND
                if notes and c <= len(notes) and notes[c - 1]:
                    cell.comment = Comment(notes[c - 1], 'Follow-up sheet')
                continue
            if v == '':
                continue
            value, fmt = cell_value(heads[c - 1] if c <= len(heads) else '', v)
            cell = ws.cell(r, c, value)
            if fmt: cell.number_format = fmt
            if isinstance(value, str) and '\n' in value:
                from openpyxl.styles import Alignment
                cell.alignment = Alignment(wrap_text=True, vertical='top')
    ws.freeze_panes = 'A2'
    for i, w in enumerate(widths or [], 1):
        ws.column_dimensions[ws.cell(1, i).column_letter].width = w
    return ws

wb = Workbook()
readme = wb.active; readme.title = 'Read me'
for i, text in enumerate([
    'Follow-up Program sheet: SAMPLE with made-up demo data (same data as the demo videos)',
    '',
    'Program plan: one line per regular session. This is what you fill in with your real sessions (one example per frequency here).',
    '   Leave "Line ID" empty for new lines: the script fills it in. Type times with AM/PM. Pick Day, Frequency and Week of month from the lists.',
    'Slots: made by the script from Program plan (Program -> Update slots now), one row per date. Do not type dates here.',
    '   Organisers may change Status (Open/Cancelled), Places, Contact and Notes. Volunteers and Still needed are written by the page.',
    '   Here: Asha Rao booked herself and Ravi Kumar (via Asha Rao) for 4 Saturdays; 24 Oct is Cancelled to show the notice.',
    'Speakers: the names volunteers can pick from (Sr. No., Speaker, Mobile). Meena Iyer has no mobile yet, to show the first-time mobile box.',
    'Update report: what the last "Update slots now" did and anything that needs you.',
    '',
    'To make a sample copy with real data: replace the rows under the titles in Program plan and Speakers with real ones',
    '   (and clear the Line ID column). Slots, Volunteers and Update report can be left as they are, or emptied.',
    'Privacy: real mobile numbers are fine in your own copy. When you share a copy with Claude, it is kept out of the code repository;',
    '   you can also replace the last digits of the numbers if you prefer.',
    'Hover over a column title to read who fills it. Hover notes, drop-down lists, date and time formats match the live sheet.']):
    readme.cell(i + 1, 1, text).font = BOLD if i == 0 else Font()
readme.column_dimensions['A'].width = 140

plan_ws = add_tab(wb, 'Program plan', data['plan'], data['planNotes'], [8, 11, 10, 10, 30, 13, 12, 20, 22, 28, 8, 13, 13, 20, 34])
slots_ws = add_tab(wb, 'Slots', data['slots'], data['slotNotes'], [14, 16, 6, 10, 10, 20, 22, 28, 20, 8, 11, 36, 8, 14, 16])
add_tab(wb, 'Speakers', data['speakers'], None, [8, 18, 14])
add_tab(wb, 'Update report', [[r[0] if r else ''] for r in data['report']] if data['report'] else [['']], None, [120])

# Drop-down lists as on the live sheet (kept on a hidden tab, because "Daily, weekdays only (Mon-Fri)" contains a comma).
lists = wb.create_sheet('Lists')
days = ['Every day'] + data['days'][1:] + data['days'][:1]
for col, values in enumerate([days, data['freqs'], ['1st', '2nd', '3rd', '4th', 'last'], ['Open', 'Cancelled']], 1):
    for r, v in enumerate(values, 1):
        lists.cell(r, col, v)
lists.sheet_state = 'hidden'
def dropdown(ws, title, col_letter, n):
    heads = [c.value for c in ws[1]]
    letter = ws.cell(1, heads.index(title) + 1).column_letter
    dv = DataValidation(type='list', formula1='=Lists!$%s$1:$%s$%d' % (col_letter, col_letter, n), allow_blank=True)
    dv.add('%s2:%s1000' % (letter, letter)); ws.add_data_validation(dv)
dropdown(plan_ws, 'Day', 'A', len(days)); dropdown(plan_ws, 'Frequency', 'B', len(data['freqs']))
dropdown(plan_ws, 'Week of month', 'C', 5); dropdown(slots_ws, 'Status', 'D', 2)

out = ROOT / 'docs' / 'followup' / 'sample-followup-sheet.xlsx'
wb.save(out)
print('wrote', out.relative_to(ROOT), '-', len(data['slots']) - 1, 'dates from', len(data['plan']) - 1, 'plan lines')
